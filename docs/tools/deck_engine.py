"""Tiny slide engine: one slide model -> .pptx (python-pptx, with speaker notes) and a PNG preview (Pillow)
drawn with the same geometry and Arial metrics, plus text-fit checks.

Slide model (dict):
  {'kind': 'title'|'section'|'bullets'|'image'|'table'|'stats'|'split'|'closing',
   'title': str, 'subtitle': str, 'bullets': [str], 'image': path, 'caption': str,
   'rows': [[...]], 'col_widths': [..], 'stats': [(value, label)], 'notes': str}
"""
import os
import textwrap

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

W_PX, H_PX = 1280, 720            # preview pixels; slide is 13.333 x 7.5 in at 96 dpi
EMU_PER_PX = 9525                 # 914400 EMU per inch / 96
FONT = 'C:/Windows/Fonts/arial.ttf'
FONT_B = 'C:/Windows/Fonts/arialbd.ttf'
GREEN_900 = (0x0F, 0x2A, 0x1A)
GREEN_700 = (0x14, 0x36, 0x1F)
LIME = (0xB5, 0xF0, 0x00)
LIME_INK = (0x3F, 0x6A, 0x00)
PAPER = (0xFA, 0xFA, 0xF7)
INK = (0x11, 0x1A, 0x14)
MUTED = (0x55, 0x60, 0x58)
LINE = (0xC9, 0xD3, 0xCB)
SOFT = (0xE8, 0xF1, 0xE9)
WHITE = (0xFF, 0xFF, 0xFF)

PROBLEMS = []


def px(v):
    return Emu(int(v * EMU_PER_PX))


def pt_to_px(pt):
    return pt * 96 / 72


def font(pt, bold=False):
    return ImageFont.truetype(FONT_B if bold else FONT, int(round(pt_to_px(pt))))


def wrap(text, pt, width_px, bold=False):
    """Greedy word wrap with Arial metrics (approximates PowerPoint)."""
    f = font(pt, bold)
    lines = []
    for para in text.split('\n'):
        words = para.split(' ')
        cur = ''
        for w in words:
            trial = (cur + ' ' + w).strip()
            if f.getlength(trial) <= width_px or not cur:
                cur = trial
            else:
                lines.append(cur)
                cur = w
        lines.append(cur)
    return lines


def line_h(pt, spacing=1.2):
    return pt_to_px(pt) * spacing


class Deck:
    def __init__(self, title, out_pptx, preview_dir):
        self.title = title
        self.out_pptx = out_pptx
        self.preview_dir = preview_dir
        os.makedirs(preview_dir, exist_ok=True)
        self.prs = Presentation()
        self.prs.slide_width = px(W_PX)
        self.prs.slide_height = px(H_PX)
        self.blank = self.prs.slide_layouts[6]
        self.n = 0

    # ---------------------------------------------------------------- primitives (both outputs)
    def rect(self, slide, img, x, y, w, h, fill, line=None):
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, px(x), px(y), px(w), px(h))
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor(*fill)
        if line:
            shp.line.color.rgb = RGBColor(*line)
            shp.line.width = Pt(0.75)
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False
        img.rectangle([x, y, x + w, y + h], fill=fill, outline=line)
        return shp

    def text(self, slide, img, x, y, w, h, text, pt, color=INK, bold=False, align='left', anchor='top', spacing=1.15, check=True, name=''):
        lines = wrap(text, pt, w - 4, bold)
        need = len(lines) * line_h(pt, spacing)
        if check and need > h + 1:
            PROBLEMS.append(f'{self.title} slide {self.n}: text overflows ({name or text[:40]!r}: needs {need:.0f}px, box {h}px)')
        tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = px(2)
        tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {'top': MSO_ANCHOR.TOP, 'middle': MSO_ANCHOR.MIDDLE, 'bottom': MSO_ANCHOR.BOTTOM}[anchor]
        for i, para in enumerate(text.split('\n')):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}[align]
            p.line_spacing = spacing
            r = p.add_run()
            r.text = para
            r.font.size = Pt(pt)
            r.font.bold = bold
            r.font.name = 'Arial'
            r.font.color.rgb = RGBColor(*color)
        # preview
        f = font(pt, bold)
        lh = line_h(pt, spacing)
        top = y if anchor == 'top' else (y + (h - need) / 2 if anchor == 'middle' else y + h - need)
        for i, ln in enumerate(lines):
            lw = f.getlength(ln)
            lx = x + 2 if align == 'left' else (x + (w - lw) / 2 if align == 'center' else x + w - lw - 2)
            img.text((lx, top + i * lh + (lh - pt_to_px(pt)) / 2), ln, font=f, fill=color)
        return need

    def bullets(self, slide, img, x, y, w, h, items, pt=20, color=INK):
        cy = y
        gap = pt_to_px(pt) * 0.55
        total = sum(len(wrap(t, pt, w - 34)) * line_h(pt) for t in items) + gap * (len(items) - 1)
        if total > h + 1:
            PROBLEMS.append(f'{self.title} slide {self.n}: bullets overflow (needs {total:.0f}px, box {h}px)')
        for t in items:
            lines = wrap(t, pt, w - 34)
            bh = len(lines) * line_h(pt)
            d = pt_to_px(pt) * 0.32
            self.rect(slide, img, x + 4, cy + line_h(pt) / 2 - d / 2, d, d, LIME_INK)
            self.text(slide, img, x + 30, cy, w - 30, bh + 2, t, pt, color, check=False)
            cy += bh + gap

    def image(self, slide, img, path, x, y, w, h, border=True, crop=None):
        im = Image.open(path).convert('RGB')
        if crop:  # (x0, y0, x1, y1) in source pixels; the cropped copy is what goes into the .pptx
            im = im.crop(crop)
            path = os.path.join(self.preview_dir, f'crop-{self.n:02d}.png')
            im.save(path)
        r = min(w / im.width, h / im.height)
        iw, ih = im.width * r, im.height * r
        ix, iy = x + (w - iw) / 2, y + (h - ih) / 2
        slide.shapes.add_picture(path, px(ix), px(iy), px(iw), px(ih))
        img_res = im.resize((max(1, int(iw)), max(1, int(ih))), Image.LANCZOS)
        img._image.paste(img_res, (int(ix), int(iy)))
        if border:
            img.rectangle([ix, iy, ix + iw, iy + ih], outline=LINE)
            line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, px(ix), px(iy), px(iw), px(ih))
            line.fill.background()
            line.line.color.rgb = RGBColor(*LINE)
            line.line.width = Pt(0.75)
            line.shadow.inherit = False
        hi_dpi = 'diagrams' in path  # diagrams are rendered at 2x
        if (r * (2 if hi_dpi else 1)) < 0.45:
            PROBLEMS.append(f'{self.title} slide {self.n}: image {os.path.basename(path)} shrunk to {r:.0%} (may be hard to read)')

    def table(self, slide, img, x, y, w, rows, col_widths=None, pt=None, max_h=None):
        cols = len(rows[0])
        col_widths = col_widths or [1] * cols
        tot = sum(col_widths)
        cw = [w * c / tot for c in col_widths]

        def measure(size):
            hs = [max(len(wrap(str(c), size, cw[i] - 16, bold=False)) for i, c in enumerate(r)) * line_h(size) + 14 for r in rows]
            return hs, sum(hs)
        if pt is None:  # largest size that fits
            for pt in (20, 18, 16, 14, 12):
                heights, total_h = measure(pt)
                if not max_h or total_h <= max_h:
                    break
        heights, total_h = measure(pt)
        if max_h and total_h > max_h + 1:
            PROBLEMS.append(f'{self.title} slide {self.n}: table too tall ({total_h:.0f}px > {max_h}px)')
        gt = slide.shapes.add_table(len(rows), cols, px(x), px(y), px(w), px(total_h)).table
        for i, c in enumerate(cw):
            gt.columns[i].width = px(c)
        cy = y
        for ri, r in enumerate(rows):
            gt.rows[ri].height = px(heights[ri])
            cx = x
            for ci, val in enumerate(r):
                cell = gt.cell(ri, ci)
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(*(SOFT if ri == 0 else WHITE))
                cell.margin_left = cell.margin_right = px(8)
                cell.margin_top = cell.margin_bottom = px(6)
                tf = cell.text_frame
                tf.word_wrap = True
                tf.paragraphs[0].text = ''
                run = tf.paragraphs[0].add_run()
                run.text = str(val)
                run.font.size = Pt(pt)
                run.font.name = 'Arial'
                run.font.bold = ri == 0
                run.font.color.rgb = RGBColor(*(GREEN_700 if ri == 0 else INK))
                img.rectangle([cx, cy, cx + cw[ci], cy + heights[ri]], fill=SOFT if ri == 0 else WHITE, outline=LINE)
                f = font(pt, ri == 0)
                for li, ln in enumerate(wrap(str(val), pt, cw[ci] - 16, bold=False)):
                    img.text((cx + 8, cy + 6 + li * line_h(pt)), ln, font=f, fill=GREEN_700 if ri == 0 else INK)
                cx += cw[ci]
            cy += heights[ri]
        return total_h

    # ---------------------------------------------------------------- slide templates
    def new(self, bg=PAPER):
        self.n += 1
        slide = self.prs.slides.add_slide(self.blank)
        im = Image.new('RGB', (W_PX, H_PX), bg)
        dr = ImageDraw.Draw(im)
        dr._image = im
        self.rect(slide, dr, 0, 0, W_PX, H_PX, bg)
        return slide, dr, im

    def header(self, slide, dr, title):
        self.rect(slide, dr, 0, 0, W_PX, 6, GREEN_700)
        self.rect(slide, dr, 56, 44, 8, 44, LIME)
        self.text(slide, dr, 78, 40, W_PX - 140, 54, title, 28, GREEN_900, bold=True, anchor='middle', name='title')

    def footer(self, slide, dr, dark=False):
        c = (0xB8, 0xC7, 0xBC) if dark else MUTED
        self.text(slide, dr, 56, H_PX - 34, 600, 20, f'Amaranth v2 · {self.title}', 10, c, check=False)
        self.text(slide, dr, W_PX - 156, H_PX - 34, 100, 20, str(self.n), 10, c, align='right', check=False)

    def notes(self, slide, text):
        if not text or len(text.split()) < 12:
            PROBLEMS.append(f'{self.title} slide {self.n}: speaker notes missing or too short')
        slide.notes_slide.notes_text_frame.text = text or ''

    def words_check(self, s):
        words = len(' '.join([s.get('title', ''), s.get('subtitle', ''), ' '.join(s.get('bullets', []))]).split())
        if s['kind'] in ('bullets', 'split', 'stack') and words > 48:
            PROBLEMS.append(f'{self.title} slide {self.n}: {words} words (aim ≤ 40)')

    def add(self, s):
        kind = s['kind']
        dark = kind in ('title', 'section', 'closing')
        slide, dr, im = self.new(GREEN_900 if dark else PAPER)
        if kind == 'title':
            self.rect(slide, dr, 0, 0, 18, H_PX, LIME)
            self.text(slide, dr, 90, 150, 1000, 30, s.get('kicker', 'AMARANTH · VISAYAS STATE UNIVERSITY'), 14, LIME, bold=True)
            self.text(slide, dr, 90, 200, 1100, 190, s['title'], 44, WHITE, bold=True, name='title')
            self.text(slide, dr, 90, 400, 1100, 90, s.get('subtitle', ''), 22, (0xD9, 0xE8, 0xB0))
            self.text(slide, dr, 90, 560, 1100, 70, s.get('meta', ''), 14, (0xB8, 0xC7, 0xBC))
        elif kind == 'section':
            self.rect(slide, dr, 90, 300, 120, 8, LIME)
            self.text(slide, dr, 90, 320, 1100, 140, s['title'], 40, WHITE, bold=True, name='title')
            self.text(slide, dr, 90, 470, 1100, 80, s.get('subtitle', ''), 20, (0xD9, 0xE8, 0xB0))
        elif kind == 'closing':
            self.text(slide, dr, 90, 220, 1100, 120, s['title'], 44, WHITE, bold=True, align='left', name='title')
            self.text(slide, dr, 90, 350, 1100, 160, s.get('subtitle', ''), 22, (0xD9, 0xE8, 0xB0))
        else:
            self.header(slide, dr, s['title'])
            top, bottom = 120, H_PX - 56
            if kind == 'bullets':
                if s.get('subtitle'):
                    self.text(slide, dr, 80, top, 1120, 34, s['subtitle'], 18, MUTED)
                    top += 48
                self.bullets(slide, dr, 80, top, 1120, bottom - top, s['bullets'], s.get('pt', 24))
            elif kind == 'image':
                cap_h = 30 if s.get('caption') else 0
                self.image(slide, dr, s['image'], 70, top, W_PX - 140, bottom - top - cap_h - 6, crop=s.get('crop'))
                if s.get('caption'):
                    self.text(slide, dr, 70, bottom - cap_h, W_PX - 140, cap_h, s['caption'], 13, MUTED, align='center')
            elif kind == 'split':
                lw = s.get('left_w', 520)
                self.bullets(slide, dr, 70, top, lw, bottom - top, s['bullets'], s.get('pt', 20))
                self.image(slide, dr, s['image'], 70 + lw + 30, top, W_PX - 140 - lw - 30, bottom - top - (30 if s.get('caption') else 0), crop=s.get('crop'))
                if s.get('caption'):
                    self.text(slide, dr, 70 + lw + 30, bottom - 26, W_PX - 140 - lw - 30, 26, s['caption'], 12, MUTED, align='center')
            elif kind == 'stack':
                ih = s.get('image_h', 250)
                self.image(slide, dr, s['image'], 70, top, W_PX - 140, ih, border=bool(s.get('crop')), crop=s.get('crop'))
                self.bullets(slide, dr, 80, top + ih + 30, 1120, bottom - top - ih - 30, s['bullets'], s.get('pt', 22))
            elif kind == 'table':
                if s.get('subtitle'):
                    self.text(slide, dr, 80, top, 1120, 30, s['subtitle'], 16, MUTED)
                    top += 40
                self.table(slide, dr, 70, top, W_PX - 140, s['rows'], s.get('col_widths'), s.get('pt'), max_h=bottom - top)
            elif kind == 'stats':
                items = s['stats']
                n = len(items)
                cols = s.get('cols', min(n, 4))
                rows_n = (n + cols - 1) // cols
                gw = (W_PX - 140 - (cols - 1) * 24) / cols
                gh = min(190, (bottom - top - (60 if s.get('subtitle') else 0) - (rows_n - 1) * 24) / rows_n)
                for i, (val, lab) in enumerate(items):
                    cx = 70 + (i % cols) * (gw + 24)
                    cy = top + (i // cols) * (gh + 24)
                    self.rect(slide, dr, cx, cy, gw, gh, WHITE, LINE)
                    self.rect(slide, dr, cx, cy, 6, gh, LIME)
                    self.text(slide, dr, cx + 22, cy + 16, gw - 36, gh * 0.48, val, 40, GREEN_900, bold=True, name='stat')
                    self.text(slide, dr, cx + 22, cy + gh * 0.55, gw - 36, gh * 0.42, lab, 15, MUTED, name='stat label')
                if s.get('subtitle'):
                    self.text(slide, dr, 70, bottom - 44, W_PX - 140, 40, s['subtitle'], 15, MUTED)
        self.footer(slide, dr, dark)
        self.notes(slide, s.get('notes', ''))
        self.words_check(s)
        im.save(os.path.join(self.preview_dir, f'{self.n:02d}.png'))

    def save(self):
        self.prs.save(self.out_pptx)
        return self.n


def contact_sheet(preview_dir, out, cols=4, scale=0.25):
    files = sorted(f for f in os.listdir(preview_dir) if f.endswith('.png') and f[:2].isdigit())
    tw, th = int(W_PX * scale), int(H_PX * scale)
    rows = (len(files) + cols - 1) // cols
    sheet = Image.new('RGB', (cols * (tw + 10) + 10, rows * (th + 10) + 10), (200, 200, 200))
    for i, f in enumerate(files):
        im = Image.open(os.path.join(preview_dir, f)).resize((tw, th), Image.LANCZOS)
        sheet.paste(im, (10 + (i % cols) * (tw + 10), 10 + (i // cols) * (th + 10)))
    sheet.save(out)
