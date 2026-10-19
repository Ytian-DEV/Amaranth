"""Style pandoc's default reference.docx: Amaranth green headings, Arial, table style, header and page-number footer."""
import sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

src, out = sys.argv[1], sys.argv[2]
d = Document(src)
GREEN = RGBColor(0x14, 0x36, 0x1F)
INK = RGBColor(0x11, 0x1A, 0x14)
LIME_INK = RGBColor(0x3F, 0x6A, 0x00)
MUTED = RGBColor(0x55, 0x60, 0x58)


def font(style, size=None, bold=None, color=None, name='Arial', italic=None):
    f = style.font
    f.name = name
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rfonts.set(qn(a), name)
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
        if rfonts.get(qn(a)) is not None:
            del rfonts.attrib[qn(a)]
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = color


styles = d.styles
for name in ['Normal', 'Body Text', 'First Paragraph', 'Compact']:
    if name in [s.name for s in styles]:
        font(styles[name], 10.5, color=INK)
        pf = styles[name].paragraph_format
        pf.space_after = Pt(6)
        pf.line_spacing = 1.15
for name, size in [('Heading 1', 18), ('Heading 2', 14), ('Heading 3', 12), ('Heading 4', 11)]:
    st = styles[name]
    font(st, size, bold=True, color=GREEN if name != 'Heading 3' else LIME_INK)
    st.paragraph_format.space_before = Pt(16 if name == 'Heading 1' else 12)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.keep_with_next = True
styles['Heading 1'].paragraph_format.page_break_before = False
font(styles['Title'], 28, bold=True, color=GREEN)
styles['Title'].paragraph_format.space_before = Pt(160)
font(styles['Subtitle'], 15, color=LIME_INK)
for name in ['Date', 'Author']:
    if name in [s.name for s in styles]:
        font(styles[name], 11, color=MUTED)
if 'Source Code' in [s.name for s in styles]:
    font(styles['Source Code'], 8.5, name='Consolas', color=INK)
for name in ['Verbatim Char']:
    if name in [s.name for s in styles]:
        font(styles[name], 9, name='Consolas', color=INK)
for name in ['Caption', 'Image Caption', 'Table Caption']:
    if name in [s.name for s in styles]:
        font(styles[name], 9, italic=True, color=MUTED)
if 'Block Text' in [s.name for s in styles]:
    font(styles['Block Text'], 10, color=MUTED)

# table style used by pandoc ("Table"): borders + header shading
tbl = styles['Table']
tblPr = tbl.element.find(qn('w:tblPr'))
if tblPr is None:
    tblPr = OxmlElement('w:tblPr')
    tbl.element.append(tblPr)
borders = OxmlElement('w:tblBorders')
for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
    e = OxmlElement(f'w:{edge}')
    e.set(qn('w:val'), 'single')
    e.set(qn('w:sz'), '4')
    e.set(qn('w:space'), '0')
    e.set(qn('w:color'), 'C9D3CB')
    borders.append(e)
old = tblPr.find(qn('w:tblBorders'))
if old is not None:
    tblPr.remove(old)
tblPr.append(borders)
mar = OxmlElement('w:tblCellMar')
for edge, w in (('left', 80), ('right', 80), ('top', 30), ('bottom', 30)):
    e = OxmlElement(f'w:{edge}')
    e.set(qn('w:w'), str(w))
    e.set(qn('w:type'), 'dxa')
    mar.append(e)
tblPr.append(mar)
# header row shading via conditional formatting
tblStylePr = OxmlElement('w:tblStylePr')
tblStylePr.set(qn('w:type'), 'firstRow')
rpr = OxmlElement('w:rPr')
b = OxmlElement('w:b')
rpr.append(b)
col = OxmlElement('w:color')
col.set(qn('w:val'), '14361F')
rpr.append(col)
tblStylePr.append(rpr)
tcpr = OxmlElement('w:tcPr')
shd = OxmlElement('w:shd')
shd.set(qn('w:val'), 'clear')
shd.set(qn('w:color'), 'auto')
shd.set(qn('w:fill'), 'E8F1E9')
tcpr.append(shd)
tblStylePr.append(tcpr)
tbl.element.append(tblStylePr)
font(tbl, 9.5, color=INK)

# page setup + header/footer with page number
for section in d.sections:
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.left_margin = section.right_margin = Cm(2.2)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.different_first_page_header_footer = True
    hp = section.header.paragraphs[0]
    hp.text = 'Amaranth v2 — Visayas State University'
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in hp.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = MUTED
        r.font.name = 'Arial'
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run('Page ')
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED
    r2 = fp.add_run()
    r2.font.size = Pt(8)
    r2.font.color.rgb = MUTED
    for kind, text in (('begin', None), ('instr', ' PAGE '), ('separate', None), ('t', '1'), ('end', None)):
        if kind == 'instr':
            el = OxmlElement('w:instrText')
            el.set(qn('xml:space'), 'preserve')
            el.text = text
        elif kind == 't':
            el = OxmlElement('w:t')
            el.text = text
        else:
            el = OxmlElement('w:fldChar')
            el.set(qn('w:fldCharType'), kind)
        r2._r.append(el)
d.save(out)
print('reference ok')
